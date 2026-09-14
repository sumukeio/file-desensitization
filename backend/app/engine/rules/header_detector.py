"""表头行启发式识别（RFC003 阶段 A）。"""
from typing import List, Tuple

from openpyxl.worksheet.worksheet import Worksheet

from backend.app.engine.rules.profiles import KNOWN_HEADERS_KEYWORDS, PSEUDO_COLUMN_MAPPING

DEFAULT_SCAN_LIMIT = 15


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _is_numeric(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _keyword_hits(text: str) -> int:
    lowered = text.lower()
    hits = 0
    for kw in KNOWN_HEADERS_KEYWORDS:
        if kw.lower() in lowered:
            hits += 1
            break
    for key in PSEUDO_COLUMN_MAPPING:
        if key in lowered:
            hits += 1
            break
    return hits


def score_row_as_header(ws: Worksheet, row: int, max_col: int, max_row: int) -> float:
    """为候选行计算“像表头”的分数，越高越可能是表头。"""
    non_empty: List[Tuple[int, object]] = []
    for col in range(1, max_col + 1):
        val = ws.cell(row, col).value
        if _cell_text(val):
            non_empty.append((col, val))

    if not non_empty:
        return -999.0

    score = 0.0
    count = len(non_empty)

    # 非空列越多，越像表头
    score += min(count, 12) * 2.0

    text_count = 0
    num_count = 0
    keyword_hits = 0
    for _, val in non_empty:
        if _is_numeric(val):
            num_count += 1
        else:
            text_count += 1
        keyword_hits += _keyword_hits(_cell_text(val))

    score += keyword_hits * 8.0
    if text_count >= num_count:
        score += text_count * 1.5

    # 标题行降权：通常只有 1 个合并说明单元格
    if count == 1:
        title_text = _cell_text(non_empty[0][1])
        score -= 15.0
        if len(title_text) > 15:
            score -= 10.0
        for hint in ("汇总", "降序", "报表", "明细", "统计"):
            if hint in title_text:
                score -= 5.0

    # 下一行更像数据行 → 当前行更像表头
    if row < max_row:
        next_vals = [ws.cell(row + 1, col).value for col in range(1, max_col + 1)]
        next_nonempty = [v for v in next_vals if _cell_text(v)]
        next_num = sum(1 for v in next_nonempty if _is_numeric(v))
        if next_nonempty and next_num >= max(1, len(next_nonempty) // 2):
            score += 12.0

    return score


def detect_header_row(ws: Worksheet, scan_limit: int = DEFAULT_SCAN_LIMIT) -> int:
    """扫描前若干行，返回最可能的表头行（1-based）。"""
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    scan_end = min(scan_limit, max_row)

    best_row = 1
    best_score = float("-inf")

    for row in range(1, scan_end + 1):
        score = score_row_as_header(ws, row, max_col, max_row)
        if score > best_score:
            best_score = score
            best_row = row

    return best_row


def resolve_header_row(
    ws: Worksheet,
    header_row_mode: str = "auto",
    manual_header_row: int | None = None,
    scan_limit: int = DEFAULT_SCAN_LIMIT,
    sheet_override: int | None = None,
) -> int:
    """按优先级解析表头行：Sheet 纠错覆盖 > 全局 manual > 自动识别。"""
    max_row = ws.max_row or 1
    if sheet_override is not None:
        return max(1, min(int(sheet_override), max_row))
    if header_row_mode == "manual" and manual_header_row is not None:
        return max(1, min(int(manual_header_row), max_row))
    return detect_header_row(ws, scan_limit=scan_limit)
