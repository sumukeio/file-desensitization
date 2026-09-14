"""RFC001：可选列名 / Sheet 名脱敏 — 引擎层单元与集成测试。"""
import openpyxl
from pathlib import Path

from backend.app.engine.excel_processor import ExcelProcessor
from backend.app.engine.rules.pseudonym import PseudonymPool
from backend.app.core.config import WORKSPACE_DIR


class TestStructureOptions:
    def setup_method(self):
        self.processor = ExcelProcessor()
        self.output_dir = WORKSPACE_DIR / "temp" / "test_output"
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.fixtures_dir = self.output_dir / "structure_fixtures"
        self.fixtures_dir.mkdir(parents=True, exist_ok=True)

    def _make_sample_workbook(self, path: Path) -> None:
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "利润明细"
        ws1["A1"] = "ASIN"
        ws1["B1"] = "店铺"
        ws1["C1"] = "销量"
        ws1["A2"] = "B082TEST001"
        ws1["B2"] = "测试店铺A"
        ws1["C2"] = 100

        ws2 = wb.create_sheet("库存明细")
        ws2["A1"] = "ASIN"
        ws2["B1"] = "所属仓库"
        ws2["A2"] = "B082TEST001"
        ws2["B2"] = "美东仓"

        wb.save(str(path))
        wb.close()

    def test_default_keeps_headers_and_sheet_names(self):
        """默认关闭：列名与 Sheet 名均不变。"""
        src = self.fixtures_dir / "default_src.xlsx"
        out = self.fixtures_dir / "default_out.xlsx"
        self._make_sample_workbook(src)

        result = self.processor.process(src, out, profile="ai_friendly")
        wb = openpyxl.load_workbook(str(out))

        assert "利润明细" in wb.sheetnames
        assert "库存明细" in wb.sheetnames
        ws = wb["利润明细"]
        assert ws.cell(1, 1).value == "ASIN"
        assert ws.cell(1, 2).value == "店铺"
        assert str(ws.cell(2, 1).value).startswith("ASIN_")
        assert not any(s.get("mask_type") == "HEADER" for s in result["diff_samples"])
        assert not any(s.get("mask_type") == "SHEET" for s in result["diff_samples"])
        wb.close()

    def test_mask_headers_only_preserves_entity_match(self):
        """仅勾选列名：表头变假名，但数据行仍按原始列名做 ASIN 假名化。"""
        src = self.fixtures_dir / "headers_src.xlsx"
        out = self.fixtures_dir / "headers_out.xlsx"
        self._make_sample_workbook(src)

        result = self.processor.process(
            src,
            out,
            profile="ai_friendly",
            custom_options={"mask_headers": True, "mask_sheet_names": False},
        )
        wb = openpyxl.load_workbook(str(out))

        assert "利润明细" in wb.sheetnames
        ws = wb["利润明细"]
        assert ws.cell(1, 1).value == "列_01"
        assert ws.cell(1, 2).value == "列_02"
        # 数据仍按原始「ASIN」列匹配假名化
        assert str(ws.cell(2, 1).value).startswith("ASIN_")
        assert str(ws.cell(2, 2).value).startswith("店铺_")
        assert isinstance(ws.cell(2, 3).value, (int, float))
        assert any(s.get("mask_type") == "HEADER" for s in result["diff_samples"])
        assert not any(s.get("mask_type") == "SHEET" for s in result["diff_samples"])
        wb.close()

    def test_mask_sheet_names_only(self):
        """仅勾选 Sheet 名：工作表重命名，列名不变。"""
        src = self.fixtures_dir / "sheets_src.xlsx"
        out = self.fixtures_dir / "sheets_out.xlsx"
        self._make_sample_workbook(src)

        result = self.processor.process(
            src,
            out,
            profile="ai_friendly",
            custom_options={"mask_headers": False, "mask_sheet_names": True},
        )
        wb = openpyxl.load_workbook(str(out))

        assert "利润明细" not in wb.sheetnames
        assert "库存明细" not in wb.sheetnames
        assert "Sheet_01" in wb.sheetnames
        assert "Sheet_02" in wb.sheetnames
        ws = wb["Sheet_01"]
        assert ws.cell(1, 1).value == "ASIN"
        assert ws.cell(1, 2).value == "店铺"
        assert any(s.get("mask_type") == "SHEET" for s in result["diff_samples"])
        assert not any(s.get("mask_type") == "HEADER" for s in result["diff_samples"])
        assert "利润明细" in result["sheet_renames"]
        assert result["sheet_renames"]["利润明细"] != "利润明细"
        wb.close()

    def test_mask_both_headers_and_sheets(self):
        """两开关都开：列名与 Sheet 名均处理。"""
        src = self.fixtures_dir / "both_src.xlsx"
        out = self.fixtures_dir / "both_out.xlsx"
        self._make_sample_workbook(src)

        result = self.processor.process(
            src,
            out,
            profile="ai_friendly",
            custom_options={"mask_headers": True, "mask_sheet_names": True},
        )
        wb = openpyxl.load_workbook(str(out))

        assert "Sheet_01" in wb.sheetnames
        ws = wb["Sheet_01"]
        assert ws.cell(1, 1).value == "列_01"
        assert str(ws.cell(2, 1).value).startswith("ASIN_")
        types = {s.get("mask_type") for s in result["diff_samples"]}
        assert "HEADER" in types
        assert "SHEET" in types
        wb.close()

    def test_cross_sheet_same_header_consistent_pseudonym(self):
        """同名列名跨 Sheet 假名一致；同 ASIN 值跨表一致。"""
        src = self.fixtures_dir / "cross_src.xlsx"
        out = self.fixtures_dir / "cross_out.xlsx"
        self._make_sample_workbook(src)
        pool = PseudonymPool()

        self.processor.process(
            src,
            out,
            profile="ai_friendly",
            custom_options={"mask_headers": True, "mask_sheet_names": False},
            shared_pseudo_pool=pool,
        )
        wb = openpyxl.load_workbook(str(out))
        ws1 = wb["利润明细"]
        ws2 = wb["库存明细"]
        # 两表都有「ASIN」列名 → 应映射为同一个「列_xx」
        assert ws1.cell(1, 1).value == ws2.cell(1, 1).value == "列_01"
        # 同一 ASIN 值跨表一致
        assert ws1.cell(2, 1).value == ws2.cell(2, 1).value
        wb.close()

    def test_unique_sheet_title_sanitizes_illegal_chars(self):
        used = set()
        name = ExcelProcessor._unique_sheet_title(r"a/b?c*[d]", used)
        assert "/" not in name
        assert "?" not in name
        assert "*" not in name
        assert "[" not in name
        assert "]" not in name
        assert len(name) <= 31

        # 撞名追加后缀
        used2 = {"Sheet_01"}
        n2 = ExcelProcessor._unique_sheet_title("Sheet_01", used2)
        assert n2 != "Sheet_01"
        assert n2 in used2

    def test_standard_mask_headers(self):
        """标准打星模式下勾选列名：表头中间打星。"""
        src = self.fixtures_dir / "std_src.xlsx"
        out = self.fixtures_dir / "std_out.xlsx"
        self._make_sample_workbook(src)

        self.processor.process(
            src,
            out,
            profile="standard_mask",
            custom_options={"mask_headers": True},
        )
        wb = openpyxl.load_workbook(str(out))
        ws = wb["利润明细"]
        assert "*" in str(ws.cell(1, 2).value)  # 「店铺」会被打星
        assert ws.cell(1, 1).value != "列_01"  # 不是假名化
        wb.close()
