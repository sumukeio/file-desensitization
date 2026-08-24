import os
import openpyxl
from pathlib import Path
from backend.app.engine.excel_processor import ExcelProcessor
from backend.app.engine.rules.pseudonym import PseudonymPool
from backend.app.core.config import WORKSPACE_DIR


class TestExampleFilesDesensitization:
    """针对 docs/Example 下真实业务数据进行的端到端脱敏验证"""

    def setup_method(self):
        self.processor = ExcelProcessor()
        self.output_dir = WORKSPACE_DIR / "temp" / "test_output"
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.shared_pool = PseudonymPool()

    def test_profit_report_desensitization(self):
        """测试 75 列高维《2601-07利润报表.xlsx》"""
        input_file = WORKSPACE_DIR / "docs" / "Example" / "2601-07利润报表.xlsx"
        assert input_file.exists()

        output_file = self.output_dir / "out_profit.xlsx"
        result = self.processor.process(
            input_path=input_file,
            output_path=output_file,
            profile="ai_friendly",
            shared_pseudo_pool=self.shared_pool,
        )

        assert result["masked_count"] > 0
        assert output_file.exists()

        # 加载脱敏后文件进行关键断言
        wb = openpyxl.load_workbook(str(output_file))
        ws = wb.active
        
        # 1. 验证表头完整性 (75列表头未被破坏)
        headers = [ws.cell(1, col).value for col in range(1, ws.max_column + 1)]
        assert "ASIN" in headers
        assert "店铺" in headers
        assert "毛利润" in headers
        assert "毛利率" in headers
        assert "销量" in headers

        # 2. 验证 ASIN / 店铺 已假名化
        asin_col = headers.index("ASIN") + 1
        shop_col = headers.index("店铺") + 1
        val_asin_r2 = ws.cell(2, asin_col).value
        val_shop_r2 = ws.cell(2, shop_col).value
        
        assert str(val_asin_r2).startswith("ASIN_")
        assert str(val_shop_r2).startswith("店铺_")

        # 3. 验证数值列保留（如销量、毛利率）
        sales_col = headers.index("销量") + 1
        val_sales = ws.cell(2, sales_col).value
        # 数值或 None 均不应变成打星字符串
        if val_sales is not None:
            assert isinstance(val_sales, (int, float))

    def test_cross_file_relational_linkage(self):
        """测试多文件跨表关联：利润表与 FBA 库存表共享同一个 ASIN 假名映射"""
        file_profit = WORKSPACE_DIR / "docs" / "Example" / "2601-07利润报表.xlsx"
        file_fba = WORKSPACE_DIR / "docs" / "Example" / "FBA仓库明细-FBA库存-示例数据.xlsx"
        file_mapping = WORKSPACE_DIR / "docs" / "Example" / "店铺团队映射表.xlsx"

        out_profit = self.output_dir / "batch_profit.xlsx"
        out_fba = self.output_dir / "batch_fba.xlsx"
        out_mapping = self.output_dir / "batch_mapping.xlsx"

        # 共享同一个批次假名池
        batch_pool = PseudonymPool()

        # 脱敏处理 3 个文件
        self.processor.process(file_profit, out_profit, "ai_friendly", shared_pseudo_pool=batch_pool)
        self.processor.process(file_fba, out_fba, "ai_friendly", shared_pseudo_pool=batch_pool)
        self.processor.process(file_mapping, out_mapping, "ai_friendly", shared_pseudo_pool=batch_pool)

        # 验证 3 个文件均成功生成
        assert out_profit.exists()
        assert out_fba.exists()
        assert out_mapping.exists()

        # 验证跨表映射池不为空且具备多个实体类别
        assert "ASIN" in batch_pool._mappings
        assert "店铺" in batch_pool._mappings
