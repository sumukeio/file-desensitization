"""上传 API 联调：结构脱敏开关透传。"""
import io

import openpyxl
from fastapi.testclient import TestClient

from backend.app.main import app


def _make_xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试表"
    ws["A1"] = "ASIN"
    ws["A2"] = "B082TEST001"
    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    return buf.getvalue()


class TestUploadStructureOptions:
    def setup_method(self):
        self.client = TestClient(app)

    def test_upload_default_structure_options_false(self):
        resp = self.client.post(
            "/api/upload",
            files={"files": ("test.xlsx", _make_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"profile": "ai_friendly"},
        )
        assert resp.status_code == 200
        tasks = resp.json()
        assert len(tasks) == 1
        assert tasks[0]["mask_headers"] is False
        assert tasks[0]["mask_sheet_names"] is False

    def test_completed_task_payload_includes_header_rows(self):
        """处理结果中的 header_rows 能进入任务模型，并出现在 Diff 接口。"""
        from pathlib import Path
        import time

        from backend.app.core.config import WORKSPACE_DIR
        from backend.app.core.queue import task_queue_manager
        from backend.app.engine.excel_processor import ExcelProcessor
        from backend.app.schemas.task import TaskStatus

        resp = self.client.post(
            "/api/upload",
            files={"files": ("test.xlsx", _make_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"profile": "ai_friendly", "header_row_mode": "auto"},
        )
        assert resp.status_code == 200
        task_id = resp.json()[0]["task_id"]
        task = task_queue_manager.get_task(task_id)
        assert task is not None

        # TestClient 同步请求不会持续驱动后台 asyncio Worker，这里直接走引擎路径验证透传
        input_path = WORKSPACE_DIR / "backend" / "temp" / "uploads" / f"{task_id}_{task.file_name}"
        output_path = WORKSPACE_DIR / "backend" / "temp" / "processed" / f"{task_id}_desensitized_{task.file_name}"
        result = ExcelProcessor().process(input_path, output_path, profile="ai_friendly")
        task.status = TaskStatus.COMPLETED
        task.progress = 100
        task.masked_count = result["masked_count"]
        task.diff_samples = result["diff_samples"]
        task.header_rows = result.get("header_rows") or {}
        task.sheet_renames = result.get("sheet_renames") or {}
        task.completed_at = time.time()
        task.output_filename = f"desensitized_{task.file_name}"

        assert task.header_rows.get("测试表") == 1

        diff = self.client.get(f"/api/tasks/{task_id}/diff")
        assert diff.status_code == 200
        body = diff.json()
        assert body["header_rows"]["测试表"] == 1
        assert body["header_row_mode"] == "auto"
        assert "header_row_overrides" in body

        tasks_resp = self.client.get("/api/tasks")
        listed = {t["task_id"]: t for t in tasks_resp.json()["tasks"]}
        assert listed[task_id]["header_rows"]["测试表"] == 1

    def test_reprocess_sets_sheet_override(self):
        """纠错重跑会写入 header_row_overrides，且 Diff 可读到。"""
        from pathlib import Path
        import time

        from backend.app.core.config import WORKSPACE_DIR, UPLOAD_DIR
        from backend.app.core.queue import task_queue_manager
        from backend.app.engine.excel_processor import ExcelProcessor
        from backend.app.schemas.task import TaskStatus

        resp = self.client.post(
            "/api/upload",
            files={"files": ("test.xlsx", _make_xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"profile": "ai_friendly", "header_row_mode": "auto"},
        )
        assert resp.status_code == 200
        task_id = resp.json()[0]["task_id"]
        task = task_queue_manager.get_task(task_id)

        input_path = UPLOAD_DIR / f"{task_id}_{task.file_name}"
        output_path = WORKSPACE_DIR / "backend" / "temp" / "processed" / f"{task_id}_desensitized_{task.file_name}"
        result = ExcelProcessor().process(input_path, output_path, profile="ai_friendly")
        task.status = TaskStatus.COMPLETED
        task.progress = 100
        task.masked_count = result["masked_count"]
        task.diff_samples = result["diff_samples"]
        task.header_rows = result.get("header_rows") or {}
        task.sheet_renames = result.get("sheet_renames") or {}
        task.completed_at = time.time()
        task.output_filename = f"desensitized_{task.file_name}"

        re_resp = self.client.post(
            f"/api/tasks/{task_id}/reprocess",
            data={"sheet_name": "测试表", "header_row": "2", "clear_override": "false"},
        )
        assert re_resp.status_code == 200
        body = re_resp.json()
        assert body["status"] == "pending"
        assert body["header_row_overrides"]["测试表"] == 2

        # 同步模拟 worker：按 overrides 再跑一遍
        result2 = ExcelProcessor().process(
            input_path,
            output_path,
            profile="ai_friendly",
            custom_options={
                "header_row_mode": "auto",
                "header_row_overrides": body["header_row_overrides"],
            },
        )
        task.status = TaskStatus.COMPLETED
        task.header_rows = result2.get("header_rows") or {}
        task.header_row_overrides = body["header_row_overrides"]
        task.diff_samples = result2["diff_samples"]
        task.progress = 100

        diff = self.client.get(f"/api/tasks/{task_id}/diff")
        assert diff.status_code == 200
        assert diff.json()["header_row_overrides"]["测试表"] == 2
        assert diff.json()["header_rows"]["测试表"] == 2
